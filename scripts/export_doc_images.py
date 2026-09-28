#!/usr/bin/env python3
import os
import csv
import json
import subprocess

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

WORKSPACE_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
CONTENT_DIR = os.path.join(WORKSPACE_ROOT, 'src', 'content')
CSV_OUTPUT = os.path.join(WORKSPACE_ROOT, 'documentation_images.csv')
XLSX_OUTPUT = os.path.join(WORKSPACE_ROOT, 'documentation_images.xlsx')
DATA_JSON_OUTPUT = os.path.join(WORKSPACE_ROOT, 'src', 'data', 'docs-images.json')
ROOT_JSON_OUTPUT = os.path.join(WORKSPACE_ROOT, 'docs-images.json')

R2_BASE_URL = 'https://pub-5672a1ed48d647e6aa3ba78e02e5445d.r2.dev/docs-images/'

def get_records_from_node():
    node_script = """
const fs = require('fs');
const path = require('path');
const { getAllDocs, getDocsTree } = require('./src/lib/markdown.js');

const allDocs = getAllDocs();
const docsMap = new Map();
allDocs.forEach(d => docsMap.set(d.id, d));

const { tree } = getDocsTree();

const orderedDocIds = [];
function traverse(nodes) {
  nodes.forEach(node => {
    orderedDocIds.push(node.id);
    if (node.children && node.children.length > 0) {
      traverse(node.children);
    }
  });
}
traverse(tree);

const records = [];

orderedDocIds.forEach(id => {
  const doc = docsMap.get(id);
  if (!doc) return;

  const lines = doc.content.split('\\n');
  let currentHeading = doc.title;

  lines.forEach((line) => {
    const headingMatch = line.match(/^(#{1,6})\\s+(.*)$/);
    if (headingMatch) {
      currentHeading = headingMatch[2].trim().replace(/\\*\\*/g, '');
    }

    const imgMatches = [...line.matchAll(/!\\[(.*?)\\]\\((.*?\\.(?:png|jpg|jpeg|gif|svg|webp))\\)/gi)];
    imgMatches.forEach(m => {
      const decodedImgRef = decodeURIComponent(m[2]);
      const dirOfDoc = path.dirname(doc.fullPath);
      let targetPath = path.join(dirOfDoc, decodedImgRef);
      if (!fs.existsSync(targetPath)) {
        const altPath = path.join(process.cwd(), 'src/content', decodedImgRef);
        if (fs.existsSync(altPath)) {
          targetPath = altPath;
        }
      }

      const relToWorkspace = path.relative(process.cwd(), targetPath).replace(/\\\\/g, '/');
      const imgName = path.basename(targetPath);

      records.push({
        websiteImagePath: 'http://localhost:3000/docs/' + doc.slug,
        titleName: currentHeading,
        imagePath: relToWorkspace,
        imageName: imgName
      });
    });
  });
});

process.stdout.write(JSON.stringify(records));
"""
    result = subprocess.run(
        ['node', '-e', node_script],
        cwd=WORKSPACE_ROOT,
        capture_output=True,
        text=True,
        check=True
    )
    return json.loads(result.stdout)

def main():
    print("Collecting documentation images and website routes...")
    records = get_records_from_node()
    print(f"Found {len(records)} images across documentation.")

    headers = ['website Image path', 'Title name', 'Image path', 'Image name', 'Image URL']

    # 1. Export CSV
    with open(CSV_OUTPUT, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in records:
            r2_url = f"{R2_BASE_URL}{r['imageName']}"
            writer.writerow([r['websiteImagePath'], r['titleName'], r['imagePath'], r['imageName'], r2_url])
    print(f"✓ Saved CSV file: {CSV_OUTPUT}")

    # 2. Export Unified JSON files
    images_map = {}
    metadata_map = {}
    items_list = []

    for r in records:
        img_name = r['imageName']
        url = f"{R2_BASE_URL}{img_name}"
        images_map[img_name] = url
        entry = {
            'imageName': img_name,
            'url': url,
            'title': r['titleName'],
            'websitePath': r['websiteImagePath'],
            'localPath': r['imagePath']
        }
        metadata_map[img_name] = entry
        items_list.append(entry)

    unified_data = {
        'baseUrl': R2_BASE_URL,
        'totalImages': len(records),
        'images': images_map,
        'metadata': metadata_map,
        'list': items_list
    }

    os.makedirs(os.path.dirname(DATA_JSON_OUTPUT), exist_ok=True)
    with open(DATA_JSON_OUTPUT, 'w', encoding='utf-8') as f:
        json.dump(unified_data, f, indent=2)
    print(f"✓ Saved Unified JSON: {DATA_JSON_OUTPUT}")

    with open(ROOT_JSON_OUTPUT, 'w', encoding='utf-8') as f:
        json.dump(unified_data, f, indent=2)
    print(f"✓ Saved Unified JSON: {ROOT_JSON_OUTPUT}")

    # 3. Export Styled Excel (XLSX)
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Documentation Images'
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        header_align = Alignment(horizontal='left', vertical='center')

        thin_side = Side(border_style='thin', color='E2E8F0')
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.row_dimensions[1].height = 28
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        regular_font = Font(name='Calibri', size=10, color='1E293B')
        link_font = Font(name='Calibri', size=10, color='1D4ED8', underline='single')
        stripe_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        cell_align = Alignment(horizontal='left', vertical='center')

        for row_idx, r in enumerate(records, 2):
            ws.row_dimensions[row_idx].height = 22
            row_fill = stripe_fill if row_idx % 2 == 0 else white_fill
            r2_url = f"{R2_BASE_URL}{r['imageName']}"

            # Col 1: website Image path
            cell1 = ws.cell(row=row_idx, column=1, value=r['websiteImagePath'])
            cell1.hyperlink = r['websiteImagePath']
            cell1.font = link_font
            cell1.fill = row_fill
            cell1.alignment = cell_align
            cell1.border = border

            # Col 2: Title name
            cell2 = ws.cell(row=row_idx, column=2, value=r['titleName'])
            cell2.font = regular_font
            cell2.fill = row_fill
            cell2.alignment = cell_align
            cell2.border = border

            # Col 3: Image path
            cell3 = ws.cell(row=row_idx, column=3, value=r['imagePath'])
            cell3.font = regular_font
            cell3.fill = row_fill
            cell3.alignment = cell_align
            cell3.border = border

            # Col 4: Image name
            cell4 = ws.cell(row=row_idx, column=4, value=r['imageName'])
            cell4.font = regular_font
            cell4.fill = row_fill
            cell4.alignment = cell_align
            cell4.border = border

            # Col 5: Cloudflare R2 URL
            cell5 = ws.cell(row=row_idx, column=5, value=r2_url)
            cell5.hyperlink = r2_url
            cell5.font = link_font
            cell5.fill = row_fill
            cell5.alignment = cell_align
            cell5.border = border

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:E{len(records) + 1}'

        wb.save(XLSX_OUTPUT)
        print(f"✓ Saved Excel file: {XLSX_OUTPUT}")

if __name__ == '__main__':
    main()
